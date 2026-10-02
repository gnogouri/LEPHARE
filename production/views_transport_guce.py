"""
Production Transport (facultés) : import des bordereaux GUCE (« ressortie de primes facultés »).

Règles d'URANUS reprises sans changement (production/exceltopostgresql.py et procédures en base) :
- période : début le 1er ou le 16, fin le 15 ou le dernier jour du mois (28, 29, 30 ou 31), même
  mois ; pas de chevauchement avec une période déjà importée (la même période exceptée) ;
- fichier : 5 Mo au plus, ligne d'en-tête détectée, exactement les 29 colonnes GUCE ;
- dates de requête et de certificat comprises dans la période ; lignes sans n° de requête ignorées ;
- souscripteur rapproché d'un client par ressemblance de nom : introuvable ou ambigu = refus (LE PHARE
  laisse l'opérateur choisir le client parmi les candidats, ou celui des imports précédents) ;
- certificats enregistrés par n° de requête (mis à jour s'il existe déjà), historique par période,
  puis sp_creation_devis_transport : un devis confirmé en contrat (offre 109, police à
  l'abonnement) par client et par police GUCE, une ligne par certificat, et la quittance.
  L'accessoire et la prime TTC sont diminués de 500 F par certificat (part AFS-CI) : c'est le
  « TOTAL GENERAL » du bordereau.

Ajouts LE PHARE (la route d'URANUS importationfichierguce reste disponible et inchangée) :
- analyse sans rien écrire : aperçu, contrôles par ligne, contrats qui seront générés ;
- import tout ou rien : URANUS enregistre chaque certificat au fil de l'eau, une erreur en cours de
  fichier laissait un import partiel ;
- refus anticipé d'une période déjà convertie en contrat (la procédure la refuse de toute façon,
  mais après avoir réécrit les certificats).
"""

import calendar
import json
import re
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import numpy as np
import pandas as pd
from django.db import connection, models, transaction
from django.http import HttpResponse
from rest_framework import permissions, status
from rest_framework.authentication import BasicAuthentication
from rest_framework.decorators import (
    api_view,
    authentication_classes,
    parser_classes,
    permission_classes,
)
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from account.models import UranusUser
from core.utils import convert_to_date
from customer.models import Client
from institutionnel.authentication import KnoxOrDemoTokenAuthentication

from .exceltopostgresql import (
    check_date_validity,
    normalize_text,
    date_columns,
    excel_column_mapping,
    excel_expected_columns,
    get_customer_id,
    locate_header_and_read_excel,
    map_columns,
    max_file_size,
)
from .models import CertificatTransport, HistoriqueImportationCertificat
from rapidfuzz import fuzz

TABLE_CERTIFICATS = "stdcertificattransport"
ACCESSOIRE_CLIENT_PAR_CERTIFICAT = Decimal("500")  # part AFS-CI, règle de sp_generation_devis_transport_client
COLONNES_MONTANTS = ["valeurassurance", "primenette", "accessoire", "taxe", "primettc", "accessoireafsci"]
AUTHENTIFICATION = [KnoxOrDemoTokenAuthentication, BasicAuthentication]


class ErreurBordereau(Exception):
    """Fichier ou période inutilisable : message affiché tel quel."""


def _nombre(valeur):
    if valeur is None:
        return 0.0
    try:
        if pd.isna(valeur):
            return 0.0
    except (TypeError, ValueError):
        pass
    try:
        return float(valeur)
    except (TypeError, ValueError):
        return 0.0


def _jour(valeur):
    if valeur is None:
        return None
    try:
        if pd.isna(valeur):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    try:
        return convert_to_date(str(valeur))
    except ValueError:
        return None


def _iso(valeur):
    j = _jour(valeur)
    return j.isoformat() if j else None


def _texte(valeur):
    """Texte d'une cellule : un numéro lu comme nombre sans « .0 », retour chariot Excel (_x000D_) retiré."""
    if valeur is None:
        return ""
    try:
        if pd.isna(valeur):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    return str(valeur).replace("_x000D_", "").strip()


def _premiere_ligne(valeur):
    return _texte(valeur).split("\n")[0].split(" / ")[0].strip()


def accessoire_client(accessoire):
    """500 F par certificat dont l'accessoire atteint 500 F (CASE WHEN accessoire < 500 THEN 0 ELSE 500)."""
    return ACCESSOIRE_CLIENT_PAR_CERTIFICAT if Decimal(str(_nombre(accessoire))) >= ACCESSOIRE_CLIENT_PAR_CERTIFICAT else Decimal("0")


# ------------------------------------------------------------------ période
def periodes_importees(exclure=None):
    """Périodes de l'historique examinées par check_date_validity (année précédente et en cours)."""
    annee = datetime.now().year
    qs = HistoriqueImportationCertificat.objects.filter(
        date_debut_periode__gte=date(annee - 1, 1, 1),
        date_debut_periode__lt=date(annee, 1, 1) + timedelta(days=365),
    ).values_list("date_debut_periode", "date_fin_periode")
    return [p for p in qs if p != exclure]


def controler_periode(debut, fin):
    """Mêmes règles que check_date_validity, avec la raison du refus."""
    erreurs = []
    if debut >= fin:
        erreurs.append("La date de fin doit suivre la date de début.")
    if (debut.year, debut.month) != (fin.year, fin.month):
        erreurs.append("La période doit tenir dans un seul mois.")
    if debut.day not in (1, 16):
        erreurs.append("La période commence le 1er ou le 16 du mois.")
    dernier_jour = calendar.monthrange(fin.year, fin.month)[1]
    if fin.day not in (15, dernier_jour):
        erreurs.append(f"La période se termine le 15 ou le {dernier_jour} (dernier jour du mois).")
    for p_debut, p_fin in periodes_importees(exclure=(debut, fin)):
        if p_debut <= fin and debut <= p_fin:
            erreurs.append(
                f"La période chevauche le bordereau déjà importé du {p_debut:%d/%m/%Y} au {p_fin:%d/%m/%Y}."
            )
    if not erreurs and not check_date_validity(debut, fin):
        erreurs.append("Période invalide.")
    return erreurs


def import_existant(debut, fin):
    """Bordereau déjà importé pour cette période et contrats qui en sont issus."""
    historique = HistoriqueImportationCertificat.objects.filter(
        date_debut_periode=debut, date_fin_periode=fin
    ).first()
    if not historique:
        return None
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT d.iddevis, d.confirme, c.numeropolice
            FROM stdimportationcertificatdevis i
            JOIN stddevis d ON d.iddevis = i.iddevis
            LEFT JOIN stdcontrat c ON c.iddevis = d.iddevis
            WHERE i.idhistoriqueimportation = %s
            """,
            [historique.id_importation],
        )
        devis = cursor.fetchall()
    return {
        "IdImportation": historique.id_importation,
        "DateImport": historique.date_creation.isoformat() if historique.date_creation else None,
        "Fichier": historique.nom_fichier_excel,
        "Succes": historique.succes,
        "Confirme": any(d[1] for d in devis),
        "Polices": sorted({d[2] for d in devis if d[2]}),
    }


# ------------------------------------------------------------------ lecture du fichier
def _periode_du_titre(fichier, nb_lignes):
    """Titre du bordereau et « Période du JJ/MM/AAAA au JJ/MM/AAAA » (lignes au-dessus de l'en-tête)."""
    fichier.seek(0)
    brut = pd.read_excel(fichier, header=None, dtype=str, nrows=max(nb_lignes, 1))
    fichier.seek(0)
    titre, periode = "", None
    for valeur in brut.fillna("").values.ravel():
        texte = " ".join(str(valeur).split())
        if not titre and "BORDEREAU" in texte.upper():
            titre = texte
        trouve = re.search(r"P[ée]riode du (\d{2}/\d{2}/\d{4}) au (\d{2}/\d{2}/\d{4})", texte, re.I)
        if trouve and not periode:
            periode = (datetime.strptime(trouve.group(1), "%d/%m/%Y").date(),
                       datetime.strptime(trouve.group(2), "%d/%m/%Y").date())
    return titre, periode


def lire_bordereau(fichier):
    if fichier.size > max_file_size:
        raise ErreurBordereau(f"La taille du fichier excède {max_file_size / (1024 * 1024):.2f} Mo.")
    try:
        fichier.seek(0)
        df, ligne_entete = locate_header_and_read_excel(
            fichier, expected_columns=excel_expected_columns, sheet_name=0, date_columns=date_columns
        )
    except Exception as erreur:
        raise ErreurBordereau(f"Erreur lors de la lecture du fichier : {erreur}")
    finally:
        fichier.seek(0)

    reelles = {str(c).strip().lower() for c in df.columns}
    attendues = {c.lower() for c in excel_expected_columns}
    if reelles != attendues:
        message = "Erreur dans les colonnes du fichier Excel."
        if reelles - attendues:
            message += " Colonnes inconnues : " + ", ".join(sorted(reelles - attendues)) + "."
        if attendues - reelles:
            message += " Colonnes manquantes : " + ", ".join(sorted(attendues - reelles)) + "."
        raise ErreurBordereau(message)

    titre, periode_titre = _periode_du_titre(fichier, ligne_entete)
    df = map_columns(df, excel_column_mapping)
    df["_ligne_excel"] = df.index + ligne_entete + 2  # numéro de ligne vu dans Excel
    df["daterequete"] = pd.to_datetime(df["daterequete"], errors="coerce", dayfirst=True)
    df["datecertificat"] = pd.to_datetime(df["datecertificat"], errors="coerce", dayfirst=True)
    return df, titre, periode_titre


# ------------------------------------------------------------------ analyse
SEUIL_RESSEMBLANCE = 85  # seuil de get_customer_id


def _candidats(nom):
    """Clients retenus par get_customer_id (mêmes mots-clés, ressemblance >= 85 %), avec leur score."""
    nom_norm = normalize_text(nom)
    requete = Client.objects.none()
    for mot in nom_norm.split()[:3]:
        requete = requete | Client.objects.filter(Nom__icontains=mot)
    candidats = []
    for id_client, nom_client, prenoms in requete.values_list("IdClient", "Nom", "Prenoms"):
        score = fuzz.token_set_ratio(nom_norm, normalize_text(nom_client))
        if score >= SEUIL_RESSEMBLANCE:
            candidats.append({"IdClient": id_client, "Nom": f"{nom_client} {prenoms or ''}".strip(), "Score": round(score)})
    return sorted(candidats, key=lambda c: -c["Score"])


def identifier_souscripteurs(noms, correspondances=None):
    """
    Client de chaque souscripteur : rapprochement automatique d'URANUS (get_customer_id), sauf choix
    explicite de l'opérateur quand le nom est introuvable ou ambigu. Le client déjà retenu pour ce
    souscripteur lors des imports précédents est proposé.
    """
    correspondances = correspondances or {}
    resultat = {}
    for nom in noms:
        automatique = get_customer_id(nom)
        with connection.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT DISTINCT t.idclienturanus, TRIM(c.nom || ' ' || COALESCE(c.prenoms, ''))
                FROM {TABLE_CERTIFICATS} t JOIN stdclient c ON c.idclient = t.idclienturanus
                WHERE t.nomsouscripteur = %s
                """,
                [nom],
            )
            precedents = [{"IdClient": i, "Nom": n} for i, n in cursor.fetchall()]
        choisi = correspondances.get(nom)
        if choisi and Client.objects.filter(pk=choisi).exists():
            id_client, statut = int(choisi), "CHOISI"
        else:
            id_client = automatique
            statut = "IDENTIFIE" if automatique > 0 else ("AMBIGU" if automatique == -1 else "INTROUVABLE")
        resultat[nom] = {
            "Souscripteur": _premiere_ligne(nom),
            "NomFichier": nom,
            "IdClient": id_client if id_client > 0 else None,
            "Statut": statut,
            "Candidats": _candidats(nom) if statut != "IDENTIFIE" else [],
            "Precedents": precedents,
            "_id": id_client,
        }
    return resultat


def _compagnies_par_assureur(assureurs):
    """Rapprochement de sp_generation_devis_transport_client : début du libellé GUCE = raison sociale."""
    resultat = {}
    with connection.cursor() as cursor:
        for assureur in assureurs:
            cursor.execute(
                """
                SELECT idcompagnie, raisonsociale FROM stdcompagnie
                WHERE SUBSTRING(%s FROM 1 FOR LENGTH(TRIM(raisonsociale))) = TRIM(raisonsociale)
                """,
                [assureur],
            )
            resultat[assureur] = cursor.fetchall()
    return resultat


def analyser(fichier, debut=None, fin=None, correspondances=None):
    df, titre, periode_titre = lire_bordereau(fichier)
    if not debut or not fin:
        if not periode_titre:
            raise ErreurBordereau("Indiquez la période du bordereau (introuvable dans le titre du fichier).")
        debut, fin = periode_titre
    erreurs_periode = controler_periode(debut, fin)
    deja = import_existant(debut, fin)
    if deja and deja["Confirme"]:
        erreurs_periode.append(
            "Ce bordereau a déjà été importé et converti en contrat"
            + (f" (police {', '.join(deja['Polices'])})" if deja["Polices"] else "")
            + " : un nouvel import de la même période est refusé."
        )
    avertissements_generaux = []
    if periode_titre and periode_titre != (debut, fin):
        avertissements_generaux.append(
            f"Le titre du fichier indique la période du {periode_titre[0]:%d/%m/%Y} au {periode_titre[1]:%d/%m/%Y}."
        )

    ignorees = int(df["numerorequete"].isna().sum())
    df = df.dropna(subset=["numerorequete"]).copy()
    df["numerorequete"] = df["numerorequete"].astype(str).str.strip()

    souscripteurs = identifier_souscripteurs(df["nomsouscripteur"].dropna().unique(), correspondances)
    clients = {nom: s["_id"] for nom, s in souscripteurs.items()}
    noms_clients = {
        i: f"{n} {p or ''}".strip()
        for i, n, p in Client.objects.filter(pk__in=[i for i in clients.values() if i > 0]).values_list("IdClient", "Nom", "Prenoms")
    }
    compagnies = _compagnies_par_assureur([_texte(a) for a in df["assureur"].dropna().unique()])

    existants = {}
    with connection.cursor() as cursor:
        cursor.execute(
            f"SELECT numerorequete, datedebutperiode, datefinperiode FROM {TABLE_CERTIFICATS} WHERE numerorequete = ANY(%s)",
            [list(df["numerorequete"].unique())],
        )
        for numero, p_debut, p_fin in cursor.fetchall():
            existants[numero] = (p_debut, p_fin)
    doublons = set(df.loc[df["numerorequete"].duplicated(keep=False), "numerorequete"])

    lignes, groupes = [], {}
    for _, r in df.iterrows():
        erreurs, avertissements = [], []
        for colonne, libelle in (("daterequete", "Date requête"), ("datecertificat", "Date certificat")):
            valeur = r[colonne]
            if pd.isna(valeur):
                erreurs.append(f"{libelle} invalide ou absente.")
            elif not debut <= valeur.date() <= fin:
                erreurs.append(f"{libelle} {valeur:%d/%m/%Y} hors de la période.")
        id_client = clients.get(r["nomsouscripteur"], 0) if not pd.isna(r["nomsouscripteur"]) else 0
        souscripteur = _premiere_ligne(r["nomsouscripteur"])
        if id_client == 0:
            erreurs.append(f"Client « {souscripteur} » introuvable : créez-le dans la Clientèle ou choisissez-le.")
        elif id_client == -1:
            erreurs.append(f"Plusieurs clients proches de « {souscripteur} » : choisissez le client du souscripteur.")
        numero = r["numerorequete"]
        if numero in doublons:
            avertissements.append("N° de requête en double dans le fichier : la dernière ligne remplace les précédentes.")
        if numero in existants and existants[numero] != (debut, fin):
            p_debut, p_fin = existants[numero]
            avertissements.append(
                f"Certificat déjà importé sur le bordereau du {p_debut:%d/%m/%Y} au {p_fin:%d/%m/%Y} : il sera rattaché à cette période."
            )
        police = _texte(r["numeropolice"])
        assureur = _texte(r["assureur"])
        compagnies_trouvees = compagnies.get(assureur, [])
        if not police:
            avertissements.append("Sans n° de police : aucun contrat ne sera généré pour ce certificat.")
        if assureur and not compagnies_trouvees:
            avertissements.append("Assureur non reconnu parmi les compagnies.")
        elif len(compagnies_trouvees) > 1:
            avertissements.append("Assureur rapproché de plusieurs compagnies.")

        accessoire = _nombre(r["accessoire"])
        ligne = {
            "Ligne": int(r["_ligne_excel"]),
            "Statut": _texte(r["statut"]),
            "NumeroRequete": numero,
            "DateRequete": _iso(r["daterequete"]),
            "ReferenceCertificat": _texte(r["referencecertificat"]),
            "DateCertificat": _iso(r["datecertificat"]),
            "NumeroPolice": police,
            "Assureur": _premiere_ligne(assureur),
            "Souscripteur": souscripteur,
            "IdClient": id_client if id_client > 0 else None,
            "Client": noms_clients.get(id_client, ""),
            "Assure": _premiere_ligne(r["assure"]),
            "MoyenTransport": _texte(r["moyentransport"]),
            "Voyage": _texte(r["voyage"]),
            "DescriptionCommerciale": _texte(r["descriptioncommerciale"]),
            "NumeroDocumentTransport": _texte(r["numerodocumenttransport"]),
            "ValeurAssurance": _nombre(r["valeurassurance"]),
            "PrimeNette": _nombre(r["primenette"]),
            "Accessoire": accessoire,
            "Taxe": _nombre(r["taxe"]),
            "PrimeTtc": _nombre(r["primettc"]),
            "AccessoireAfsCi": _nombre(r["accessoireafsci"]),
            "Erreurs": erreurs,
            "Avertissements": avertissements,
        }
        lignes.append(ligne)
        if police and id_client > 0:
            cle = (id_client, police)
            g = groupes.setdefault(cle, {
                "IdClient": id_client, "Client": noms_clients.get(id_client, souscripteur), "NumeroPolice": police,
                "Assureur": _premiere_ligne(assureur),
                "Compagnie": compagnies_trouvees[0][1] if len(compagnies_trouvees) == 1 else "",
                "Certificats": 0, "ValeurAssurance": 0.0, "PrimeNette": 0.0, "Accessoire": 0.0, "Taxe": 0.0,
                "PrimeTtc": 0.0, "AccessoireClient": 0.0,
            })
            g["Certificats"] += 1
            for cle_montant in ("ValeurAssurance", "PrimeNette", "Accessoire", "Taxe", "PrimeTtc"):
                g[cle_montant] += ligne[cle_montant]
            g["AccessoireClient"] += float(accessoire_client(accessoire))

    for g in groupes.values():
        g["AccessoireNet"] = g["Accessoire"] - g["AccessoireClient"]
        g["PrimeTtcNette"] = g["PrimeTtc"] - g["AccessoireClient"]

    somme = lambda cle: sum(l[cle] for l in lignes)  # noqa: E731
    acc_client = float(sum(accessoire_client(l["Accessoire"]) for l in lignes))
    return {
        "Titre": titre,
        "PeriodeFichier": [periode_titre[0].isoformat(), periode_titre[1].isoformat()] if periode_titre else None,
        "DebutPeriode": debut.isoformat(),
        "FinPeriode": fin.isoformat(),
        "ErreursPeriode": erreurs_periode,
        "Avertissements": avertissements_generaux,
        "ImportExistant": deja,
        "Souscripteurs": [{k: v for k, v in s.items() if not k.startswith("_")} for s in souscripteurs.values()],
        "LignesIgnorees": ignorees,
        "Lignes": lignes,
        "Contrats": list(groupes.values()),
        "Totaux": {
            "Certificats": len(lignes),
            "Valides": sum(1 for l in lignes if not l["Erreurs"]),
            "EnErreur": sum(1 for l in lignes if l["Erreurs"]),
            "AvecAvertissement": sum(1 for l in lignes if l["Avertissements"]),
            "Doublons": len(doublons),
            "DejaEnBase": sum(1 for n in df["numerorequete"].unique() if n in existants),
            "ValeurAssurance": somme("ValeurAssurance"),
            "PrimeNette": somme("PrimeNette"),
            "Accessoire": somme("Accessoire"),
            "Taxe": somme("Taxe"),
            "PrimeTtc": somme("PrimeTtc"),
            "AccessoireAfsCi": somme("AccessoireAfsCi"),
            "AccessoireClient": acc_client,
            "TotalGeneral": somme("PrimeTtc") - acc_client,
        },
        "Importable": not erreurs_periode and bool(lignes) and all(not l["Erreurs"] for l in lignes),
        "_df": df,
        "_clients": clients,
    }


def _periode_demandee(donnees):
    debut, fin = donnees.get("debut_periode"), donnees.get("fin_periode")
    if not debut or not fin:
        return None, None
    try:
        return convert_to_date(debut), convert_to_date(fin)
    except ValueError:
        raise ErreurBordereau("Mauvais format de date !")


def _correspondances(donnees):
    """Choix de l'opérateur : {nom du souscripteur dans le fichier: IdClient} (JSON)."""
    brut = donnees.get("correspondances")
    if not brut:
        return {}
    try:
        valeurs = json.loads(brut) if isinstance(brut, str) else dict(brut)
        return {str(k): int(v) for k, v in valeurs.items() if v}
    except (ValueError, TypeError, AttributeError):
        raise ErreurBordereau("Choix des clients illisible.")


def _publique(analyse):
    return {k: v for k, v in analyse.items() if not k.startswith("_")}


@api_view(["POST"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def analyser_bordereau(request):
    """POST /api/transport/guce/analyse/ (fichier_excel, debut_periode et fin_periode facultatives)"""
    fichier = request.FILES.get("fichier_excel")
    if not fichier:
        return Response({"error": "Sélectionnez le fichier Excel du bordereau GUCE."}, status=status.HTTP_400_BAD_REQUEST)
    try:
        debut, fin = _periode_demandee(request.data)
        return Response(_publique(analyser(fichier, debut, fin, _correspondances(request.data))))
    except ErreurBordereau as erreur:
        return Response({"error": str(erreur)}, status=status.HTTP_400_BAD_REQUEST)


# ------------------------------------------------------------------ import
def _valeur_sql(colonne, valeur):
    if valeur is None:
        return None
    try:
        if pd.isna(valeur):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(valeur, pd.Timestamp):
        return valeur.to_pydatetime().date()
    if isinstance(valeur, np.integer):
        return int(valeur)
    if isinstance(valeur, np.floating):
        return float(valeur)
    return valeur


def _enregistrer_certificats(cursor, df, debut, fin, clients):
    """
    Même enregistrement qu'upsert_data_psycopg3 (par n° de requête), sans validation ligne à ligne.
    Champs texte : une cellule vide reste vide et un numéro lu comme nombre n'a pas de « .0 »
    (la troncature d'URANUS, faite après astype(str), écrivait « nan » et « 260018578.0 »).
    """
    longueurs = {
        champ.db_column: champ.max_length
        for champ in CertificatTransport._meta.fields
        if isinstance(champ, models.CharField)
    }
    colonnes = list(excel_column_mapping.values())
    for _, r in df.iterrows():
        ligne = {"datedebutperiode": debut, "datefinperiode": fin}
        for colonne in colonnes:
            valeur = _valeur_sql(colonne, r[colonne])
            if colonne in longueurs and valeur is not None:
                if isinstance(valeur, float) and valeur.is_integer():
                    valeur = int(valeur)
                valeur = str(valeur)[: longueurs[colonne]]
            ligne[colonne] = valeur
        ligne["idclienturanus"] = clients[r["nomsouscripteur"]]
        cursor.execute(f"SELECT 1 FROM {TABLE_CERTIFICATS} WHERE numerorequete = %s", [ligne["numerorequete"]])
        if cursor.fetchone():
            autres = [c for c in ligne if c != "numerorequete"]
            cursor.execute(
                f"UPDATE {TABLE_CERTIFICATS} SET " + ", ".join(f"{c} = %s" for c in autres) + " WHERE numerorequete = %s",
                [ligne[c] for c in autres] + [ligne["numerorequete"]],
            )
        else:
            cursor.execute(
                f"INSERT INTO {TABLE_CERTIFICATS} ({', '.join(ligne)}) VALUES ({', '.join(['%s'] * len(ligne))})",
                list(ligne.values()),
            )


def _contrats_de_l_import(id_importation):
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT d.iddevis, d.numerodevis, d.numeropoliceconnexe, d.idclient,
                   TRIM(cl.nom || ' ' || COALESCE(cl.prenoms, '')), co.raisonsociale,
                   d.primenette, d.accessoire, d.taxe, d.primettc, d.confirme,
                   c.idcontrat, c.numeropolice, q.numeroquittance,
                   (SELECT COUNT(*) FROM stddevisdetail dd WHERE dd.iddevis = d.iddevis)
            FROM stdimportationcertificatdevis i
            JOIN stddevis d ON d.iddevis = i.iddevis
            LEFT JOIN stdclient cl ON cl.idclient = d.idclient
            LEFT JOIN stdcompagnie co ON co.idcompagnie = d.idcompagnie
            LEFT JOIN stdcontrat c ON c.iddevis = d.iddevis
            LEFT JOIN stdquittance q ON q.idquittance = c.idquittance
            WHERE i.idhistoriqueimportation = %s
            ORDER BY d.iddevis
            """,
            [id_importation],
        )
        return [
            {
                "IdDevis": l[0], "NumeroDevis": l[1], "NumeroPoliceGuce": l[2], "IdClient": l[3], "Client": l[4],
                "Compagnie": l[5] or "", "PrimeNette": _nombre(l[6]), "Accessoire": _nombre(l[7]),
                "Taxe": _nombre(l[8]), "PrimeTtc": _nombre(l[9]), "Confirme": bool(l[10]), "IdContrat": l[11],
                "NumeroPolice": l[12] or "", "NumeroQuittance": l[13] or "", "Certificats": l[14],
            }
            for l in cursor.fetchall()
        ]


@api_view(["POST"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def importer_bordereau(request):
    """POST /api/transport/guce/import/ : certificats, historique, devis, contrats et quittances, tout ou rien."""
    fichier = request.FILES.get("fichier_excel")
    if not fichier:
        return Response({"error": "Sélectionnez le fichier Excel du bordereau GUCE."}, status=status.HTTP_400_BAD_REQUEST)
    try:
        debut, fin = _periode_demandee(request.data)
        analyse = analyser(fichier, debut, fin, _correspondances(request.data))
    except ErreurBordereau as erreur:
        return Response({"error": str(erreur)}, status=status.HTTP_400_BAD_REQUEST)
    if not analyse["Importable"]:
        message = (analyse["ErreursPeriode"] or ["Le bordereau contient des lignes en erreur : corrigez-les avant l'import."])[0]
        if not analyse["Lignes"]:
            message = "Aucun certificat dans le fichier."
        return Response({"error": message, "analyse": _publique(analyse)}, status=status.HTTP_400_BAD_REQUEST)

    debut, fin = convert_to_date(analyse["DebutPeriode"]), convert_to_date(analyse["FinPeriode"])
    try:
        with transaction.atomic():
            with connection.cursor() as cursor:
                _enregistrer_certificats(cursor, analyse["_df"], debut, fin, analyse["_clients"])
            historique = HistoriqueImportationCertificat.objects.filter(
                date_debut_periode=debut, date_fin_periode=fin
            ).first() or HistoriqueImportationCertificat(date_debut_periode=debut, date_fin_periode=fin)
            historique.nom_fichier_excel = str(fichier.name)[:255]
            historique.sha256_hash = None
            historique.operateur = UranusUser.objects.filter(pk=request.user.id).first()
            historique.save()
            with connection.cursor() as cursor:
                cursor.execute(
                    "CALL sp_creation_devis_transport(%s, %s, %s, %s, %s, %s);",
                    [historique.id_importation, debut, fin, request.user.id, 0, ""],
                )
                _, message = cursor.fetchone()
            contrats = _contrats_de_l_import(historique.id_importation)
    except Exception as erreur:
        return Response(
            {"error": str(erreur).split("\n")[0], "analyse": _publique(analyse)},
            status=status.HTTP_400_BAD_REQUEST,
        )
    return Response(
        {
            "Message": message or "Importation réalisée avec succès.",
            "IdImportation": historique.id_importation,
            "DebutPeriode": analyse["DebutPeriode"],
            "FinPeriode": analyse["FinPeriode"],
            "Totaux": analyse["Totaux"],
            "LignesIgnorees": analyse["LignesIgnorees"],
            "Contrats": contrats,
        },
        status=status.HTTP_201_CREATED,
    )


# ------------------------------------------------------------------ consultation
@api_view(["GET"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def liste_bordereaux(request):
    """GET /api/transport/guce/bordereaux/ : bordereaux importés, totaux et contrats générés."""
    with connection.cursor() as cursor:
        cursor.execute(
            f"""
            SELECT h.idimportation, h.datedebutperiode, h.datefinperiode, h.nomfichierexcel, h.date_creation,
                   h.succes, u.name, u.email,
                   COUNT(t.idcertificat), COALESCE(SUM(t.valeurassurance), 0), COALESCE(SUM(t.primenette), 0),
                   COALESCE(SUM(t.accessoire), 0), COALESCE(SUM(t.taxe), 0), COALESCE(SUM(t.primettc), 0),
                   COALESCE(SUM(CASE WHEN t.accessoire < 500 THEN 0 ELSE 500 END), 0),
                   STRING_AGG(DISTINCT TRIM(cl.nom || ' ' || COALESCE(cl.prenoms, '')), ', '),
                   STRING_AGG(DISTINCT TRIM(SPLIT_PART(REPLACE(t.nomsouscripteur, CHR(13), ''), CHR(10), 1)), ', ')
            FROM stdhistoriqueimportationcertificat h
            LEFT JOIN account_uranususer u ON u.id = h.idoperateur
            LEFT JOIN {TABLE_CERTIFICATS} t ON t.datedebutperiode = h.datedebutperiode AND t.datefinperiode = h.datefinperiode
            LEFT JOIN stdclient cl ON cl.idclient = t.idclienturanus
            GROUP BY h.idimportation, u.name, u.email
            ORDER BY h.datedebutperiode DESC
            """
        )
        lignes = cursor.fetchall()
    resultat = []
    for l in lignes:
        resultat.append({
            "IdImportation": l[0], "DebutPeriode": _iso(l[1]), "FinPeriode": _iso(l[2]), "Fichier": l[3] or "",
            "DateImport": l[4].isoformat() if l[4] else None, "Succes": bool(l[5]), "Operateur": l[6] or l[7] or "",
            "Certificats": l[8], "ValeurAssurance": _nombre(l[9]), "PrimeNette": _nombre(l[10]),
            "Accessoire": _nombre(l[11]), "Taxe": _nombre(l[12]), "PrimeTtc": _nombre(l[13]),
            "AccessoireClient": _nombre(l[14]), "TotalGeneral": _nombre(l[13]) - _nombre(l[14]),
            # Souscripteur du fichier GUCE ; client rattaché à l'import (les imports de 2024 pointent vers
            # des n° de client réattribués depuis la reprise de données)
            "Clients": l[15] or "", "Souscripteurs": l[16] or "", "Contrats": _contrats_de_l_import(l[0]),
        })
    return Response(resultat)


def _periode_et_filtres(request):
    """Certificats d'un bordereau (idimportation), d'un devis ou d'un contrat Transport."""
    id_importation = request.query_params.get("idimportation")
    id_devis = request.query_params.get("iddevis")
    id_contrat = request.query_params.get("idcontrat")
    with connection.cursor() as cursor:
        if id_importation:
            cursor.execute(
                "SELECT datedebutperiode, datefinperiode FROM stdhistoriqueimportationcertificat WHERE idimportation = %s",
                [id_importation],
            )
            ligne = cursor.fetchone()
            return (ligne[0], ligne[1], None, None) if ligne else None
        if id_contrat:
            cursor.execute("SELECT iddevis FROM stdcontrat WHERE idcontrat = %s AND idproduit = 6", [id_contrat])
            ligne = cursor.fetchone()
            id_devis = ligne[0] if ligne else None
        if id_devis:
            cursor.execute(
                "SELECT dateeffet::date, dateexpiration::date, idclient, numeropoliceconnexe FROM stddevis WHERE iddevis = %s AND idproduit = 6",
                [id_devis],
            )
            ligne = cursor.fetchone()
            return tuple(ligne) if ligne else None
    return None


@api_view(["GET"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def certificats_transport(request):
    """GET /api/transport/certificats/?idimportation=|iddevis=|idcontrat="""
    filtres = _periode_et_filtres(request)
    if not filtres:
        return Response({"error": "Bordereau, devis ou contrat Transport introuvable."}, status=status.HTTP_404_NOT_FOUND)
    debut, fin, id_client, police = filtres
    qs = CertificatTransport.objects.filter(date_debut_periode=debut, date_fin_periode=fin)
    if id_client:
        qs = qs.filter(id_client_uranus=id_client)
    if police:
        qs = qs.filter(numero_police=police)
    certificats = [
        {
            "NumeroRequete": c.numero_requete, "DateRequete": _iso(c.date_requete),
            "ReferenceCertificat": c.reference_certificat, "DateCertificat": _iso(c.date_certificat),
            "NumeroPolice": c.numero_police, "Assureur": _premiere_ligne(c.assureur),
            "Souscripteur": _premiere_ligne(c.nom_souscripteur), "Assure": _premiere_ligne(c.assure),
            "MoyenTransport": c.moyen_transport, "DateDebutVoyage": _iso(c.date_debut_voyage), "Voyage": c.voyage,
            "DescriptionCommerciale": _texte(c.description_commerciale), "NumeroDocumentTransport": c.numero_document_transport,
            "ValeurAssurance": _nombre(c.valeur_assurance), "PrimeNette": _nombre(c.prime_nette),
            "Accessoire": _nombre(c.accessoire), "Taxe": _nombre(c.taxe), "PrimeTtc": _nombre(c.prime_ttc),
            "AccessoireAfsCi": _nombre(c.accessoire_afs_ci),
        }
        for c in qs.order_by("date_certificat", "numero_requete")
    ]
    acc_client = float(sum(accessoire_client(c["Accessoire"]) for c in certificats))
    return Response({
        "DebutPeriode": _iso(debut), "FinPeriode": _iso(fin), "Certificats": certificats,
        "Totaux": {
            "Certificats": len(certificats),
            **{cle: sum(c[cle] for c in certificats) for cle in ("ValeurAssurance", "PrimeNette", "Accessoire", "Taxe", "PrimeTtc", "AccessoireAfsCi")},
            "AccessoireClient": acc_client,
            "TotalGeneral": sum(c["PrimeTtc"] for c in certificats) - acc_client,
        },
    })


@api_view(["GET"])
@authentication_classes(AUTHENTIFICATION)
@permission_classes([permissions.IsAuthenticated])
def exporter_bordereau(request, id_importation):
    """GET /api/transport/guce/bordereaux/<id>/excel/ : bordereau au format de la ressortie GUCE (29 colonnes)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    historique = HistoriqueImportationCertificat.objects.filter(pk=id_importation).first()
    if not historique:
        return Response({"error": "Bordereau introuvable."}, status=status.HTTP_404_NOT_FOUND)
    debut, fin = historique.date_debut_periode, historique.date_fin_periode
    certificats = list(
        CertificatTransport.objects.filter(date_debut_periode=debut, date_fin_periode=fin).order_by("-date_certificat", "-numero_requete")
    )
    champs = {f.db_column: f.attname for f in CertificatTransport._meta.fields}
    souscripteurs = sorted({_premiere_ligne(c.nom_souscripteur) for c in certificats if c.nom_souscripteur})
    # N° du bordereau dans l'année : une quinzaine = un bordereau (N°01 = 1er au 15 janvier, N°03 = 1er au 15 février)
    numero = (debut.month - 1) * 2 + (1 if debut.day == 1 else 2)

    classeur = Workbook()
    feuille = classeur.worksheets[0]
    feuille.title = "LISTE DES CERTIFICATS D'ASSURAN"
    feuille.cell(row=1, column=1, value=(
        f"BORDEREAU N°{numero:02d} DE RESSORTIE DE PRIMES FACULTES {debut.year}    - GUCE P/C {', '.join(souscripteurs)}"
    )).font = Font(bold=True, size=12)
    feuille.cell(row=3, column=1, value=f"Période du {debut:%d/%m/%Y} au {fin:%d/%m/%Y}").font = Font(bold=True)
    ligne_entete = 6
    for i, libelle in enumerate(excel_expected_columns, start=1):
        cellule = feuille.cell(row=ligne_entete, column=i, value=libelle)
        cellule.font = Font(bold=True)
        cellule.alignment = Alignment(wrap_text=True, vertical="center")
    for n, certificat in enumerate(certificats, start=ligne_entete + 1):
        for i, libelle in enumerate(excel_expected_columns, start=1):
            valeur = getattr(certificat, champs[excel_column_mapping[libelle]])
            if isinstance(valeur, Decimal):
                valeur = float(valeur)
            elif isinstance(valeur, date):
                valeur = valeur.strftime("%d/%m/%Y")
            feuille.cell(row=n, column=i, value=valeur)
    total = ligne_entete + len(certificats) + 1
    feuille.cell(row=total, column=1, value="TOTAL").font = Font(bold=True)
    colonnes_montants = {excel_column_mapping[l]: i for i, l in enumerate(excel_expected_columns, start=1)}
    for colonne in COLONNES_MONTANTS:
        somme = sum(_nombre(getattr(c, champs[colonne])) for c in certificats)
        feuille.cell(row=total, column=colonnes_montants[colonne], value=round(somme, 2)).font = Font(bold=True)
    net = sum(_nombre(c.prime_ttc) for c in certificats) - float(sum(accessoire_client(c.accessoire) for c in certificats))
    feuille.cell(row=total + 1, column=1, value="TOTAL GENERAL").font = Font(bold=True)
    feuille.cell(row=total + 1, column=colonnes_montants["valeurassurance"], value=round(net)).font = Font(bold=True)
    for i in range(1, len(excel_expected_columns) + 1):
        feuille.column_dimensions[get_column_letter(i)].width = 18

    tampon = BytesIO()
    classeur.save(tampon)
    nom = f"RESSORTIE DE PRIME FACULTES DU {debut:%d-%m-%Y} AU {fin:%d-%m-%Y}.xlsx"
    reponse = HttpResponse(
        tampon.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    reponse["Content-Disposition"] = f'attachment; filename="{nom}"'
    return reponse
