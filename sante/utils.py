import os
from collections import OrderedDict
from datetime import datetime
from decimal import Decimal
from itertools import chain

from django.conf import settings
from django.db import connection
from openpyxl import load_workbook

from core.date_parser import date_emission_du_jour
from core.utils import convert_to_date
from production.models import ComplementDevisDetailSante, DataInsertionResult, DevisDetail

from .models import (
    Adherent,
    AdherentSanteInsertionResult,
    Affilie,
    AffilieFn,
    AffilieSanteInsertionResult,
    FilialeSanteInsertionResult,
    SaisieDevisSanteEnCours,
)
from .serializers import (
    AdherentSanteInsertionSerializer,
    AffilieSanteInsertionSerializer,
)

importation_col_list = [
    "Nom",
    "Prenom",
    "NumeroCNI",
    "DateNaissance",
    "NumeroCMU",
    "Sexe",
    "NumeroTelephone",
    "Lien",
    "DateDebutConso",
    "DateAdhesion",
    "AdresseAdherent",
    "EmailAdherent",
    "NombreAffections",
    "LibelleAffections",
    "Matricule",
]


def message_erreur_base(error):
    """
    Message d'une erreur de procédure à afficher : certaines procédures Santé lèvent un
    texte de plusieurs lignes commençant par « Got exception: », dont la vraie cause est
    sur la ligne « message: ». Sinon, la première ligne.
    """
    lignes = str(error).splitlines() or [""]
    for ligne in lignes:
        ligne = ligne.strip()
        if ligne.startswith("message:"):
            return ligne[len("message:"):].strip()
    return lignes[0]


def remove_unwanted_keys(data):
    for key in data.keys():
        if key not in importation_col_list:
            del data[key]

    return data


def load_people_data_from_excel(filename, index=0):
    dict_list = []
    # Les indices des lignes et des colonnes commencent par 1
    book = load_workbook(filename, data_only=True, keep_vba=False)
    sheet = book.worksheets[index]

    # Lecture de la ligne d'entête
    column_count = sheet.max_column

    keys = [
        str(sheet.cell(row=1, column=col_index).value).lower()
        for col_index in range(1, column_count + 1)
    ]

    # Read other lines from the second one
    row_count = sheet.max_row

    for row_index in range(2, row_count + 1):
        # empty = all(isinstance(cell, EmptyCell) for cell in row)
        d = OrderedDict(
            {
                keys[col_index - 1]: (
                    "NA"
                    if str(sheet.cell(row_index, col_index).value) == "None"
                    else str(sheet.cell(row_index, col_index).value).strip()
                )
                for col_index in range(1, column_count + 1)
            }
        )

        # Ignore blank lines
        if (d["nom"] in ("NA", "N/A", "N-A")) and (
            d["datenaissance"] in ("NA", "N/A", "N-A")
        ):
            continue

        # Remove unwanted data, especially blank colums
        # d = remove_unwanted_keys(d)

        # Keep only the date not the time
        d["datenaissance"] = d["datenaissance"][:10]
        d["dateeffet"] = d["dateeffet"][:10]
        d["dateentree"] = d["dateentree"][:10]

        dict_list.append(d)
    return dict_list


# Import insured people from Excel File
def import_insured(filename, user_id, id_devis, date_effet, id_filiale):
    messages = []
    error_count = 0
    try:
        request_data_list = load_people_data_from_excel(
            filename=filename, index=0
        )

        if request_data_list is None or len(request_data_list) == 0:
            error_count += 1
            messages.append(
                {
                    "code_erreur": "-1",
                    "message": "Le fichier Excel ne contient pas de données valides",
                }
            )
            return (error_count, messages)

        # Parcourir les lignes extraites du fichier Excel en vue d'enregistrer les données
        id_adherent = 0
        id_devis = int(id_devis)
        id_filiale = int(id_filiale)
        # filiale_sante = FilialeSante.objects.filter(Q(devis=id_devis)).first()
        # if filiale_sante:
        #     filiale = filiale_sante.idfiliale

        for item in request_data_list:
            current_request_data = {}
            err = False
            qryset = None
            if str(item["lienparente"]).strip() == "A":
                current_request_data["idadherent"] = 0
            else:
                if id_adherent == -1:
                    continue
                current_request_data["idaffilie"] = 0
                current_request_data["adherent"] = id_adherent
                current_request_data["handicape"] = False

            current_request_data["lien"] = str(item["lienparente"]).strip()
            current_request_data["filiale"] = id_filiale
            current_request_data["nom"] = str(item["nom"]).strip()
            current_request_data["prenom"] = str(item["prenoms"]).strip()
            current_request_data["sexe"] = str(item["sexe"]).strip()
            current_request_data["cni"] = str(item["numerocni"]).strip()
            current_request_data["vip"] = True
            current_request_data["adresse"] = str(
                item["adresseadherent"]
            ).strip()
            current_request_data["mobile1"] = str(
                item["numerotelephone"]
            ).strip()
            current_request_data["mobile2"] = ""
            current_request_data["email"] = str(item["emailadherent"]).strip()
            current_request_data["numerocmu"] = str(item["numerocmu"]).strip()
            current_request_data["devis"] = id_devis
            current_request_data["dateeffet"] = str(
                item["dateentree"]
            )  # Ce choix n'est pas intuitif
            current_request_data["datedebutconsommation"] = str(
                item["dateeffet"]
            )
            current_request_data["datenaissance"] = str(item["datenaissance"])
            current_request_data["matricule"] = str(item["matricule"]).strip()
            current_request_data["surprimeappliquee"] = False
            current_request_data["montantsurprime"] = 0
            current_request_data["nombrepathologie"] = 0
            item["nombreaffection"] = str(item["nombreaffection"]).strip()
            if item["nombreaffection"].isdigit():
                current_request_data["nombrepathologie"] = int(
                    item["nombreaffection"]
                )

            current_request_data["groupesanguin"] = item["groupesanguin"]

            if current_request_data["lien"] == "A":
                (err, qryset) = enregistrer_adherent_sante(
                    user_id, current_request_data
                )
                message = AdherentSanteInsertionSerializer(
                    qryset, many=True
                ).data
                if not err:
                    id_adherent = message[0]["idadherent"]
                else:
                    id_adherent = -1
            else:
                (err, qryset) = enregistrer_affilie_sante(
                    user_id, current_request_data
                )
                message = AffilieSanteInsertionSerializer(
                    qryset, many=True
                ).data
            # Implementation future
            # (err, qryset) = enregistrer_adherent_sante(user_id, fichier_piece, request.POST)
            if err:
                error_count += 1

            messages.append(message)

    except Exception as error:
        print(error)
        message = {
            "statut": "-1000",
            "message": str(error),
        }
        error_count += 1
        messages.append(message)

    return (error_count, messages)


# Save quotation Santé
def save_quotation_sante(user_id, input_data):
    sql_output = None
    error_occurred = False
    IdIntermediaire = int(input_data["IdIntermediaire"])
    IdCompagnie = int(input_data["IdCompagnie"])
    IdProduit = int(input_data["IdProduit"])
    IdOffre = int(input_data["IdOffre"])
    IdAvenant = int(input_data["IdAvenant"])
    IdClient = int(input_data["IdClient"])
    IdAssure = int(input_data["IdAssure"])
    Flotte = bool(input_data["Flotte"])
    Coassurance = bool(input_data["Coassurance"])
    # convert_to_date
    DateEffet = convert_to_date(input_data["DateEffet"])
    DateExpiration = convert_to_date(input_data["DateExpiration"])
    DateEmission = convert_to_date(input_data["DateEmission"])
    # DateEffet = datetime.strptime(input_data["DateEffet"], "%d-%m-%Y").date()
    # DateExpiration = datetime.strptime(input_data["DateExpiration"], "%d-%m-%Y").date()
    # DateEmission = datetime.strptime(input_data["DateEmission"], "%d-%m-%Y").date()
    IdTarif = int(input_data["IdTarif"])
    IdDuree = 1
    if "IdDuree" in input_data:
        if input_data["IdDuree"]:
            IdDuree = int(input_data["IdDuree"])
    (
        PrimeFamille,
        PrimeAffilie,
        PrimeGlobale,
        MontantSurprime,
        MontantAccessoireManuel,
        TauxReductionCommerciale,
        TypeContrat,
        GestionnaireSante,
    ) = (
        0,
        0,
        0,
        0,
        0,
        0,
        1,
        "",
    )
    if "PrimeFamille" in input_data:
        if input_data["PrimeFamille"]:
            PrimeFamille = Decimal(input_data["PrimeFamille"])
    if "PrimeAffilie" in input_data:
        if input_data["PrimeAffilie"]:
            PrimeAffilie = Decimal(input_data["PrimeAffilie"])
    if "PrimeGlobale" in input_data:
        if input_data["PrimeGlobale"]:
            PrimeGlobale = Decimal(input_data["PrimeGlobale"])
    if "MontantSuprime" in input_data:
        if input_data["MontantSuprime"]:
            MontantSurprime = Decimal(input_data["MontantSuprime"])
    if "MontantAccessoireManuel" in input_data:
        if input_data["MontantAccessoireManuel"]:
            MontantAccessoireManuel = Decimal(
                input_data["MontantAccessoireManuel"]
            )
    # Lecture du taux d'ajustement envoyé par le frontend sous le nom "TauxReduction"
    # Convention : positif = réduction, négatif = majoration
    TauxReduction = Decimal("0")
    if "TauxReduction" in input_data and input_data["TauxReduction"] not in (None, "", 0, "0"):
        TauxReduction = Decimal(str(input_data["TauxReduction"]))
    elif "TauxReductionCommerciale" in input_data and input_data["TauxReductionCommerciale"] not in (None, "", 0, "0"):
        TauxReduction = Decimal(str(input_data["TauxReductionCommerciale"]))

    # Valeur transmise à la procédure stockée : signée (positif=réduction, négatif=majoration)
    # Le SP applique : prime * (1 - taux/100) → fonctionne pour réduction ET majoration
    TauxReductionCommerciale = TauxReduction

    if "TypeContrat" in input_data:
        if input_data["TypeContrat"]:
            TypeContrat = int(input_data["TypeContrat"])
    if "GestionnaireSante" in input_data:
        if input_data["GestionnaireSante"]:
            GestionnaireSante = str(input_data["GestionnaireSante"])
    IdDevis = int(input_data["IdDevis"])

    NumeroPoliceCompagnie = ""
    if "NumeroPoliceCompagnie" in input_data:
        NumeroPoliceCompagnie = (
            str(input_data["NumeroPoliceCompagnie"])
            if input_data["NumeroPoliceCompagnie"]
            else ""
        )

    OutputMessage = ""
    data_insertion_result_list = []
    queryset_vide = DataInsertionResult.objects.none()
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "CALL sp_creation_devis_sante(%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);",
                (
                    IdIntermediaire,
                    IdCompagnie,
                    IdProduit,
                    IdOffre,
                    IdAvenant,
                    IdClient,
                    IdAssure,
                    Flotte,
                    Coassurance,
                    DateEffet,
                    DateExpiration,
                    DateEmission,
                    IdTarif,
                    PrimeFamille,
                    PrimeAffilie,
                    PrimeGlobale,
                    MontantSurprime,
                    MontantAccessoireManuel,
                    TauxReductionCommerciale,
                    TypeContrat,
                    GestionnaireSante,
                    IdDuree,
                    NumeroPoliceCompagnie,
                    user_id,
                    IdDevis,
                    OutputMessage,
                ),
            )
            connection.commit()
            row = cursor.fetchone()
            sql_output = DataInsertionResult(
                ObjectId=row[0], OutputMessage=row[1]
            )
            data_insertion_result_list.append(sql_output)
    except Exception as error:
        error_occurred = True
        print(error)
        msg = message_erreur_base(error)
        sql_output = DataInsertionResult(ObjectId=IdDevis, OutputMessage=msg)
        data_insertion_result_list.append(sql_output)
    finally:
        if connection:
            cursor.close()
            connection.close()

    # Patch ORM : force la mise à jour de taux_reduction_commerciale après fermeture du curseur raw.
    # Le SP peut ignorer ce champ lorsque le complément existe déjà (ex. renouvellement),
    # donc on s'assure que la valeur envoyée est bien persistée.
    if not error_occurred and sql_output and sql_output.ObjectId and TauxReductionCommerciale != Decimal("0"):
        try:
            det_ids = list(
                DevisDetail.objects.filter(iddevis=sql_output.ObjectId).values_list(
                    "iddevisdetail", flat=True
                )
            )
            if det_ids:
                ComplementDevisDetailSante.objects.filter(
                    devis_detail__in=det_ids
                ).update(taux_reduction_commerciale=TauxReductionCommerciale)
        except Exception as patch_error:
            print(f"ORM patch taux_reduction_commerciale: {patch_error}")

    return (
        error_occurred,
        list(chain(queryset_vide, data_insertion_result_list)),
    )


def enregistrer_adherent_sante(user_id, input_data, fichier_piece=None):
    sql_output = None
    error_occurred = False
    # print(type(input_data))
    # print("Données de l'adhérent:",input_data)

    try:
        idadherent = int(input_data["idadherent"])
        filiale = int(input_data["filiale"])
        nom = str(input_data["nom"])
        prenom = str(input_data["prenom"])
        sexe = str(input_data["sexe"])
        cni = ""
        if "cni" in input_data:
            if input_data["cni"]:
                cni = str(input_data["cni"])
        vip = str(input_data["vip"])
        adresse = ""
        if "adresse" in input_data:
            if input_data["adresse"]:
                adresse = str(input_data["adresse"])
        mobile1 = str(input_data["mobile1"])
        mobile2 = str(input_data["mobile2"])
        email = str(input_data["email"])
        devis = int(input_data["devis"])
        # convert_to_date
        dateeffet = None
        if str(input_data["dateeffet"]) not in ("NA", "N/A", "N-A"):
            dateeffet = convert_to_date(input_data["dateeffet"])

        datedebutconsommation = None
        if str(input_data["datedebutconsommation"]) not in (
            "NA",
            "N/A",
            "N-A",
        ):
            datedebutconsommation = convert_to_date(
                input_data["datedebutconsommation"]
            )

        datenaissance = None
        if "datenaissance" in input_data:
            if input_data["datenaissance"] and input_data[
                "datenaissance"
            ] not in (
                "NA",
                "N/A",
                "N-A",
            ):
                datenaissance = convert_to_date(input_data["datenaissance"])

        matricule = ""
        if "matricule" in input_data:
            if input_data["matricule"]:
                matricule = str(input_data["matricule"])

        surprimeappliquee = False
        if "surprimeappliquee" in input_data:
            if input_data["surprimeappliquee"]:
                surprimeappliquee = bool(input_data["surprimeappliquee"])

        montantsurprime = 0
        if "montantsurprime" in input_data:
            if input_data["montantsurprime"]:
                montantsurprime = Decimal(input_data["montantsurprime"])

        nombrepathologie = 0
        if "nombrepathologie" in input_data:
            if input_data["nombrepathologie"]:
                nombrepathologie = int(input_data["nombrepathologie"])

        # Donner la possibilité d'utiliser également nombrepathologies (au pluriel) comme champ de l'objet JSON
        if not nombrepathologie and "nombrepathologies" in input_data:
            if input_data["nombrepathologies"]:
                nombrepathologie = int(input_data["nombrepathologies"])

        numerocmu = ""
        if "numerocmu" in input_data:
            if input_data["numerocmu"]:
                numerocmu = str(input_data["numerocmu"])

        groupesanguin = ""
        if "groupesanguin" in input_data:
            if input_data["groupesanguin"]:
                groupesanguin = str(input_data["groupesanguin"])

        outputmessage = ""
        data_insertion_result_list = []
        queryset_vide = AdherentSanteInsertionResult.objects.none()

        with connection.cursor() as cursor:
            cursor.execute(
                "CALL sp_saisie_adherent_sante(%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);",
                (
                    nom,
                    prenom,
                    sexe,
                    cni,
                    vip,
                    adresse,
                    mobile1,
                    mobile2,
                    email,
                    datenaissance,
                    dateeffet,
                    filiale,
                    datedebutconsommation,
                    matricule,
                    surprimeappliquee,
                    montantsurprime,
                    nombrepathologie,
                    numerocmu,
                    groupesanguin,
                    user_id,
                    idadherent,
                    devis,
                    outputmessage,
                ),
            )
            connection.commit()
            row = cursor.fetchone()
            sql_output = AdherentSanteInsertionResult(
                idadherent=row[0], devis=row[1], outputmessage=row[2]
            )
            data_insertion_result_list.append(sql_output)
            # Gestion du fichier joint
            if int(row[0]) and fichier_piece:
                adherent = Adherent.objects.get(pk=int(row[0]))
                old_path = None
                # old_path = adherent.fichier_piece.path
                if adherent.fichier_piece:
                    old_path = adherent.fichier_piece.path
                if old_path and os.path.isfile(old_path):
                    adherent.fichier_piece.name = (
                        "uploads/adherents/" + fichier_piece.name
                    )
                    new_path = (
                        settings.MEDIA_ROOT + adherent.fichier_piece.name
                    )
                    os.rename(old_path, new_path)
                else:
                    adherent.fichier_piece = fichier_piece
                adherent.save()
    except Exception as error:
        error_occurred = True
        print(error)
        msg = message_erreur_base(error)
        sql_output = AdherentSanteInsertionResult(
            idadherent=idadherent, devis=devis, outputmessage=msg
        )
        data_insertion_result_list.append(sql_output)
    finally:
        if connection:
            cursor.close()
            connection.close()
    return (
        error_occurred,
        list(chain(queryset_vide, data_insertion_result_list)),
    )


#########################################################################
#
# Enregistrer les affiliés
def enregistrer_affilie_sante(userid, input_data, fichier_piece=None):
    sql_output = None
    error_occurred = False
    # print(type(input_data))
    # print("Données de l'affilié:",input_data)

    try:
        idaffilie = int(input_data["idaffilie"])
        adherent = int(input_data["adherent"])
        lien = str(input_data["lien"])
        nom = str(input_data["nom"])
        prenom = str(input_data["prenom"])
        # cni = str(input_data["cni"])
        cni = ""
        if "cni" in input_data:
            if input_data["cni"]:
                cni = str(input_data["cni"])
        # convert_to_date
        datenaissance = None
        if str(input_data["datenaissance"]) not in ("NA", "N/A", "N-A"):
            datenaissance = convert_to_date(input_data["datenaissance"])

        # datenaissance = datetime.strptime(input_data["datenaissance"], "%d-%m-%Y").date()
        certificat = False
        if "certificat" in input_data:
            if input_data["certificat"]:
                certificat = bool(input_data["certificat"])

        matricule = ""
        if "matricule" in input_data:
            if input_data["matricule"]:
                matricule = str(input_data["matricule"])

        # observations = str(input_data["observations"])
        if "observations" in input_data:
            observations = (
                str(input_data["observations"])
                if input_data["observations"]
                else ""
            )
        else:
            observations = "RAS"

        handicape = False
        if "handicape" in input_data:
            if input_data["handicape"]:
                handicape = bool(input_data["handicape"])

        sexe = str(input_data["sexe"])
        devis = int(input_data["devis"])
        # convert_to_date
        dateeffet = None
        if str(input_data["dateeffet"]) not in ("NA", "N/A", "N-A"):
            dateeffet = convert_to_date(input_data["dateeffet"])

        datedebutconsommation = None
        if str(input_data["datedebutconsommation"]) not in (
            "NA",
            "N/A",
            "N-A",
        ):
            datedebutconsommation = convert_to_date(
                input_data["datedebutconsommation"]
            )

        surprimeappliquee = bool(input_data["surprimeappliquee"])
        montantsurprime = Decimal(input_data["montantsurprime"])

        nombrepathologie = 0
        if "nombrepathologie" in input_data:
            if input_data["nombrepathologie"]:
                nombrepathologie = int(input_data["nombrepathologie"])
        # Donner la possibilité d'utiliser également nombrepathologies (au pluriel) comme champ de l'objet JSON
        if not nombrepathologie and "nombrepathologies" in input_data:
            if input_data["nombrepathologies"]:
                nombrepathologie = int(input_data["nombrepathologies"])

        numerocmu = ""
        if "numerocmu" in input_data:
            if input_data["numerocmu"]:
                numerocmu = str(input_data["numerocmu"])

        groupesanguin = ""
        if "groupesanguin" in input_data:
            if input_data["groupesanguin"]:
                groupesanguin = str(input_data["groupesanguin"])

        mobile1 = ""
        mobile2 = ""
        if lien.upper() in ("A", "C"):
            if "mobile1" in input_data:
                if input_data["mobile1"]:
                    mobile1 = str(input_data["mobile1"])
            if "mobile2" in input_data:
                if input_data["mobile2"]:
                    mobile2 = str(input_data["mobile2"])
        outputmessage = ""
        data_insertion_result_list = []
        queryset_vide = AffilieSanteInsertionResult.objects.none()

        with connection.cursor() as cursor:
            cursor.execute(
                "CALL sp_saisie_affilie_sante(%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);",
                (
                    nom,
                    prenom,
                    datenaissance,
                    sexe,
                    mobile1,
                    mobile2,
                    cni,
                    lien,
                    dateeffet,
                    certificat,
                    matricule,
                    observations,
                    handicape,
                    datedebutconsommation,
                    surprimeappliquee,
                    montantsurprime,
                    nombrepathologie,
                    numerocmu,
                    groupesanguin,
                    userid,
                    idaffilie,
                    adherent,
                    devis,
                    outputmessage,
                ),
            )
            connection.commit()
            row = cursor.fetchone()
            sql_output = AffilieSanteInsertionResult(
                idaffilie=row[0],
                adherent=row[1],
                devis=row[2],
                outputmessage=row[3],
            )
            data_insertion_result_list.append(sql_output)
            # Gestion du fichier joint
            if int(row[0]) and fichier_piece:
                affilie = Affilie.objects.get(pk=int(row[0]))
                old_path = None
                # old_path = adherent.fichier_piece.path
                if affilie.fichier_piece:
                    old_path = affilie.fichier_piece.path
                if old_path and os.path.isfile(old_path):
                    affilie.fichier_piece.name = (
                        "uploads/affilies/" + fichier_piece.name
                    )
                    new_path = settings.MEDIA_ROOT + affilie.fichier_piece.name
                    os.rename(old_path, new_path)
                else:
                    affilie.fichier_piece = fichier_piece
                affilie.save()
    except Exception as error:
        error_occurred = True
        print(error)
        msg = message_erreur_base(error)
        sql_output = AffilieSanteInsertionResult(
            idaffilie=idaffilie,
            adherent=adherent,
            devis=devis,
            outputmessage=msg,
        )
        data_insertion_result_list.append(sql_output)
    finally:
        if connection:
            cursor.close()
            connection.close()
    return (
        error_occurred,
        list(chain(queryset_vide, data_insertion_result_list)),
    )


######################################################################################
# Enregistrer les filiales


def enregistrer_filiale_sante(userid, input_data):
    sql_output = None
    error_occurred = False
    idfiliale = int(input_data["idfiliale"])
    devis = int(input_data["devis"])
    college = int(input_data["college"])
    offresante = int(input_data["offresante"])
    zonecouverture = int(input_data["zonecouverture"])
    date_emission = date_emission_du_jour()
    date_effet = datetime.strptime(input_data["date_effet"], "%d-%m-%Y").date()
    date_expiration = datetime.strptime(
        input_data["date_expiration"], "%d-%m-%Y"
    ).date()
    source = str(input_data["source"])

    outputmessage = ""
    data_insertion_result_list = []
    queryset_vide = FilialeSanteInsertionResult.objects.none()
    try:
        # status = 0
        with connection.cursor() as cursor:
            cursor.execute(
                "CALL sp_saisie_filiale_sante(%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);",
                (
                    college,
                    offresante,
                    zonecouverture,
                    date_effet,
                    date_expiration,
                    date_emission,
                    source,
                    userid,
                    idfiliale,
                    devis,
                    outputmessage,
                ),
            )
            connection.commit()
            row = cursor.fetchone()
            sql_output = FilialeSanteInsertionResult(
                idfiliale=row[0], devis=row[1], outputmessage=row[2]
            )
            data_insertion_result_list.append(sql_output)
    except Exception as error:
        error_occurred = True
        print(error)
        msg = message_erreur_base(error)
        sql_output = FilialeSanteInsertionResult(
            idfiliale=idfiliale,
            devis=devis,
            outputmessage=msg,
        )
        data_insertion_result_list.append(sql_output)
    finally:
        if connection:
            cursor.close()
            connection.close()
    return (
        error_occurred,
        list(chain(queryset_vide, data_insertion_result_list)),
    )


def get_quotation_id(userid):
    msg = ""
    id_devis = 0
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "CALL sp_initialisation_devis_sante(%s, %s, %s);",
                (
                    userid,
                    id_devis,
                    msg,
                ),
            )
            connection.commit()
            row = cursor.fetchone()
            id_devis = row[0]
            msg = row[1]
    except Exception as error:
        print(error)
        id_devis = 0
        msg = message_erreur_base(error)
    finally:
        if connection:
            cursor.close()
            connection.close()
    return (msg, id_devis)


def get_saisie_sante_en_cours(user_id):
    msg = ""
    res = SaisieDevisSanteEnCours.objects.none()
    saisie_list = []
    try:
        with connection.cursor() as cursor:
            cursor.callproc(
                "fn_saisie_sante_en_cours",
                [
                    user_id,
                ],
            )

            result = cursor.fetchall()
            for row in result:
                dev = SaisieDevisSanteEnCours(
                    devis=row[0],
                    identifiant_devis=row[1],
                    utilisateur=row[2],
                )
                saisie_list.append(dev)
                # print(dev)
    except Exception as error:
        print(error)
        msg = message_erreur_base(error)
    else:
        if len(saisie_list) > 0:
            res = list(chain(res, saisie_list))

    finally:
        if connection:
            cursor.close()
            connection.close()

    return (msg, res)


def get_liste_affilie_sante(iddevis):
    msg = ""
    res = AffilieFn.objects.none()
    affilie_sante_list = []
    try:
        with connection.cursor() as cursor:
            cursor.callproc(
                "fn_liste_affilie_sante",
                [
                    iddevis,
                ],
            )

            result = cursor.fetchall()
            for row in result:
                aff = AffilieFn(
                    iddevis=row[0],
                    idaffilie=row[1],
                    idadherent=row[2],
                    nom=row[3],
                    prenom=row[4],
                    lien=row[5],
                    lienparente=row[6],
                    datenaissance=row[7],
                    sexe=row[8],
                    numerocni=row[9],
                    numerocmu=row[10],
                    groupesanguin=row[11],
                )
                affilie_sante_list.append(aff)
                # print(qp)
    except Exception as error:
        print(error)
        msg = str(error)
    else:
        if len(affilie_sante_list) > 0:
            res = list(chain(res, affilie_sante_list))

    finally:
        if connection:
            cursor.close()
            connection.close()

    return (msg, res)
